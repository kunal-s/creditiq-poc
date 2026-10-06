import csv
import datetime as dt
import io

from openpyxl import Workbook

from ingestion.pipeline import Ingestor
from ingestion.store import Store

from .conftest import spec, values

ROWS = [
    ["Continental Commercial Bank"],
    ["Statement of Account"],
    ["Account Name", "Test Traders Private Limited", "Account Number", "123456789012"],
    ["Branch / IFSC", "Pune / HDFC0001234", "Currency", "INR"],
    ["Statement", "01 Apr 2026 to 30 Jun 2026"],
    [],
    ["Date", "Particulars", "Reference", "Debit (INR)", "Credit (INR)", "Balance (INR)"],
    ["01 Apr 2026", "OPENING BALANCE", None, None, None, 100000],
    ["03 Apr 2026", "NEFT CR ACME", "NEFT/1", None, 50000, 150000],
    ["09 Apr 2026", "RTGS DR VENDOR", "RTGS/2", 20000, None, 130000],
    ["20 Jun 2026", "NEFT CR ACME", "NEFT/3", None, 10000, 140000],
    [],
    ["Statement Summary"],
    ["Opening Balance (01 Apr 2026)", 100000],
    ["Closing Balance (30 Jun 2026)", 140000],
]


def xlsx(rows=ROWS) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_xlsx_bank_statement_reads_cell_by_cell(cfg):
    data = xlsx()
    r = Ingestor(cfg, Store(), None, None).process_file("C", spec("stmt.xlsx", data, "application/vnd.ms-excel"), data)
    assert r.source["kind"] == "xlsx" and r.pages[0].source == "sheet" and r.pages[0].sheet == "Statement" and r.grade == "A"
    assert r.classification.type == "bank_statement" and r.classification.tier == "signals"
    d = r.documents[0]
    v = values(d)
    assert v["account_number"] == "123456789012" and v["ifsc"] == "HDFC0001234" and v["branch"] == "Pune"
    assert v["statement_period_start"] == "2026-04-01" and v["opening_balance"] == 100000 and v["closing_balance"] == 140000
    rows = [x.cells for x in d.tables["transactions"].rows]
    assert [x["balance"] for x in rows] == [100000, 150000, 130000, 140000]
    assert [x["credit"] for x in rows] == [None, 50000, None, 10000] and [x["debit"] for x in rows] == [None, None, 20000, None]
    assert next(c for c in d.validation if c.id == "running_balance").outcome == "pass"


def test_spreadsheet_evidence_is_a_cell_not_a_page(cfg):
    data = xlsx()
    d = Ingestor(cfg, Store(), None, None).process_file("C", spec("s.xlsx", data), data).documents[0]
    src = d.fields["account_number"].source
    assert src.kind == "cell" and src.sheet == "Statement" and src.range == "D3" and src.bbox is None
    row_src = d.tables["transactions"].rows[1].source
    assert row_src["sheet"] == "Statement" and row_src["range"].startswith("A9")


def test_real_excel_dates_are_read(cfg):
    rows = [list(r) for r in ROWS]
    for r, d in zip(range(7, 11), [dt.date(2026, 4, 1), dt.date(2026, 4, 3), dt.date(2026, 4, 9), dt.date(2026, 6, 20)]):
        rows[r][0] = d
    data = xlsx(rows)
    d = Ingestor(cfg, Store(), None, None).process_file("C", spec("s.xlsx", data), data).documents[0]
    assert [x.cells["date"] for x in d.tables["transactions"].rows] == ["2026-04-01", "2026-04-03", "2026-04-09", "2026-06-20"]


def test_csv_is_read_the_same_way(cfg):
    b = io.StringIO()
    csv.writer(b).writerows([r if r else [] for r in ROWS])
    data = b.getvalue().encode()
    r = Ingestor(cfg, Store(), None, None).process_file("C", spec("s.csv", data, "text/csv"), data)
    assert r.source["kind"] == "csv" and r.classification.type == "bank_statement"
    assert len(r.documents[0].tables["transactions"].rows) == 4


def test_a_corrupt_workbook_is_rejected(cfg):
    data = b"PK\x03\x04" + b"garbage" * 20
    r = Ingestor(cfg, Store(), None, None).process_file("C", spec("s.xlsx", data), data)
    assert r.status == "rejected" and r.reject.code == "corrupt"
