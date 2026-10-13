"""The Word writer (python-docx): an editable A4 file with the same text as the PDF, every character
as written. The running footer counts pages with Word's PAGE and NUMPAGES fields; credit writes in
the boxes and on the lines left blank."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime

from docx import Document as new_document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Length, Mm, Pt, RGBColor
from docx.table import Table as DocxTable, _Cell, _Row
from docx.text.paragraph import Paragraph as DocxParagraph

from . import layout

FONT = "Arial"
WIDTH = Cm(17)
INDENT = Cm(0.65)
LINE = 20
"""Points between ruled lines left to write on, as in the PDF."""
BORDER = 0.5
"""Points: a row's border, which Word adds to the row's height."""
PADDING = (50, 80, 70, 80)
"""A cell's padding (top, left, bottom, right) in twentieths of a point, as in the PDF."""
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
RULE = "A7B0BA"
HEAD_FILL = "E6EBF1"
LABEL_FILL = "F3F5F8"
EDGES = ("top", "left", "bottom", "right", "insideH", "insideV")
WHOLE_ROW = 30
"""Lines: a table row up to this tall goes to the next page whole rather than breaking across it (Word
would hide the overflow of an unbreakable row taller than a page)."""
LIMIT = 255
"""Characters: the longest a Word file's title, subject or other property may be."""
APP = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
       '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
       'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
       '<Application>CreditIQ</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>'
       '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>'
       '</Properties>').encode()
"""The file's application properties: written by CreditIQ (the template names another application)."""

# Paragraph styles: size, bold, italic, muted, space before and after (pt), keep with next, and the
# least distance between lines (pt), the PDF's leading.
STYLES = {
    "Memo title": (17, True, False, False, 0, 6, False, 21),
    "Memo note": (9, False, True, True, 0, 4, False, 12),
    "Memo subheading": (10, True, False, False, 6, 3, True, 13),
    "Memo label": (9, True, False, False, 4, 2, True, 12),
    "Memo question": (10, True, False, False, 10, 2, True, 13),
    "Memo detail": (8.5, False, False, True, 0, 3, False, 11),
    "Table caption": (9, True, False, False, 4, 3, True, 12),
    "Table note": (8.5, False, True, True, 3, 6, False, 11),
    "Table text": (layout.TABLE_SIZE, False, False, False, 0, 0, False, layout.LEADING),
    "Title label": (9, False, False, True, 0, 0, False, 12),
    "Title value": (9, True, False, False, 0, 0, False, 12),
}
SIZE = {"Normal": layout.BODY_SIZE, "Heading 1": 12, **{name: spec[0] for name, spec in STYLES.items()}}
BEFORE = {"Normal": 0, "Heading 1": 12, **{name: spec[4] for name, spec in STYLES.items()}}
BLOCK_STYLES = {"body": "Normal", "note": "Memo note", "subheading": "Memo subheading", "label": "Memo label",
                "question": "Memo question", "meta": "Memo detail"}
# What may follow each element in a table's, a row's, a cell's and a paragraph's properties (ECMA-376
# order): an element added later goes before the first of these present, where Word expects it.
AFTER_TBL_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
                     "w:tblPrChange")
AFTER_TBL_IND = ("w:tblBorders", *AFTER_TBL_BORDERS)
AFTER_TBL_CELL_MAR = ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
AFTER_TBL_HEADER = ("w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange")
AFTER_CANT_SPLIT = ("w:trHeight", "w:tblHeader", *AFTER_TBL_HEADER)
AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers",
             "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
AFTER_P_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
               "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
               "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
               "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
               "w:rPr", "w:sectPr", "w:pPrChange")
HYPHEN = re.compile(r"(?<=\w)-(?=\w)")


def _fonts(rpr) -> None:
    """One font throughout: the theme's fonts would otherwise win over a style's."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for name in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(name), None)
    for name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(name), FONT)


def _leading(style, points: float) -> None:
    """Lines at least so far apart, as the PDF sets them (Word's own spacing follows the font's)."""
    style.paragraph_format.line_spacing = Pt(points)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST


def _styles(doc) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")) if defaults is not None else None
    if rpr is not None:
        _fonts(rpr)
    normal = doc.styles["Normal"]
    _fonts(normal.element.get_or_add_rPr())
    normal.font.size = Pt(SIZE["Normal"])
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    _leading(normal, 12.5)
    normal.paragraph_format.widow_control = True
    heading = doc.styles["Heading 1"]
    _fonts(heading.element.get_or_add_rPr())
    heading.font.size = Pt(SIZE["Heading 1"])
    heading.font.bold = True
    heading.font.color.rgb = INK
    heading.paragraph_format.space_before = Pt(BEFORE["Heading 1"])
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True
    _leading(heading, 15)
    for name, size in (("Header", 8), ("Footer", 8)):
        style = doc.styles[name]
        _fonts(style.element.get_or_add_rPr())
        style.font.size = Pt(size)
        style.font.color.rgb = MUTED
    for name, (size, bold, italic, muted, before, after, keep, leading) in STYLES.items():
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = MUTED if muted else INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = keep
        _leading(style, leading)
    question = doc.styles["Memo question"].paragraph_format
    question.left_indent = INDENT
    question.first_line_indent = -INDENT


def _settings(doc) -> None:
    """Word's current layout rules (compatibility mode 15), not Word 2010's, which the template asks for."""
    for setting in doc.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")


def _unbroken(text: str, at: int) -> bool:
    """Whether the hyphen at `at` joins one reference, which a line must not break: a code such as
    XV-000004-12, GSTR-3B or BBG-2026-000004, a year such as FY 2025-26, a range of sources such as 13-23."""
    if text[at - 1].isdigit() or text[at + 1].isdigit():
        return True
    start = max(text.rfind(c, 0, at) for c in "  ([") + 1
    return text[start:at].isupper()


def _text(run, text: str) -> None:
    """The run's text: a line break as Word's, and a hyphen inside a reference as one that does not
    break the line."""
    r = run._r
    for k, line in enumerate(text.split("\n")):
        if k:
            r.add_br()
        start = 0
        for m in HYPHEN.finditer(line):
            if _unbroken(line, m.start()):
                if m.start() > start:
                    r.add_t(line[start:m.start()])
                r.append(OxmlElement("w:noBreakHyphen"))
                start = m.end()
        if start < len(line):
            r.add_t(line[start:])


def _runs(paragraph: DocxParagraph, spans: layout.Spans, size: float) -> DocxParagraph:
    for s in spans:
        run = paragraph.add_run()
        _text(run, s.text)
        if s.bold:
            run.bold = True
        if s.italic:
            run.italic = True
        if s.cite:
            run.font.size = Pt(size - 1.5)
            run.font.color.rgb = MUTED
    return paragraph


def _border(element, tag: str, edges: tuple[str, ...], present: tuple[str, ...], after: tuple[str, ...]) -> None:
    borders = OxmlElement(tag)
    for edge in edges:
        line = OxmlElement(f"w:{edge}")
        if edge in present:
            for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", RULE)):
                line.set(qn(key), value)
        else:
            line.set(qn("w:val"), "nil")
        borders.append(line)
    element.insert_element_before(borders, *after)


def _shade(cell: _Cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    for key, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shd.set(qn(key), value)
    cell._tc.get_or_add_tcPr().insert_element_before(shd, *AFTER_SHD)


def _repeat(row: _Row) -> None:
    """The header row, repeated at the top of each page the table runs onto."""
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().insert_element_before(header, *AFTER_TBL_HEADER)


def _whole(row: _Row) -> None:
    """The row moves to the next page whole, as in the PDF, rather than breaking across the page."""
    row._tr.get_or_add_trPr().insert_element_before(OxmlElement("w:cantSplit"), *AFTER_CANT_SPLIT)


def _least(row: _Row, points: float, padding: tuple[int, int, int, int] = PADDING) -> None:
    """The row at least this tall, in points, as the PDF draws it: Word adds the cells' top and bottom
    padding, and the border, to the height a row is given."""
    row.height = Pt(max(points - (padding[0] + padding[2]) / 20 - BORDER, 1))
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _twips(length: int) -> str:
    return str(Length(length).twips)


def _layout(table: DocxTable, widths: tuple[float, ...], edges: tuple[str, ...], *, indent: int = 0,
            padding: tuple[int, int, int, int] = PADDING) -> None:
    """Fixed column widths across the text width, light borders on the given edges and cell padding
    (top, left, bottom, right, in twentieths of a point)."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    total = Emu(WIDTH - indent)
    width = tbl_pr.find(qn("w:tblW"))
    if width is not None:
        width.set(qn("w:type"), "dxa")
        width.set(qn("w:w"), _twips(total))
    if indent:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), _twips(indent))
        ind.set(qn("w:type"), "dxa")
        tbl_pr.insert_element_before(ind, *AFTER_TBL_IND)
    _border(tbl_pr, "w:tblBorders", EDGES, edges, AFTER_TBL_BORDERS)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in zip(("top", "left", "bottom", "right"), padding):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    tbl_pr.insert_element_before(margins, *AFTER_TBL_CELL_MAR)
    columns = [Emu(int(total * w)) for w in widths]
    for column, w in zip(table.columns, columns):
        column.width = w
    for row in table.rows:
        for cell, w in zip(row.cells, columns):
            cell.width = w


def _field(paragraph: DocxParagraph, instruction: str) -> None:
    """A Word field (PAGE, NUMPAGES), which Word fills in as it lays out the pages."""
    def mark(kind: str) -> None:
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(char)

    mark("begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    paragraph.add_run()._r.append(code)
    mark("separate")
    paragraph.add_run("1")
    mark("end")


def _rule_below(paragraph: DocxParagraph) -> None:
    _border(paragraph._p.get_or_add_pPr(), "w:pBdr", ("top", "left", "bottom", "right"), ("bottom",), AFTER_P_BDR)


def _keep(rows: list[list[DocxParagraph]]) -> None:
    for row in rows:
        for p in row:
            p.paragraph_format.keep_with_next = True


class _Writer:
    """Adds the blocks to the document. Each block returns the paragraphs it made, row by row (a table
    row's cells together), so that a group can keep every row but its last with the next."""

    def __init__(self, doc) -> None:
        self.doc = doc
        self.after_table = self.after_space = False
        # Styles by id: python-docx looks a style's name up among all the styles each time it is set.
        self.ids = {s.name: s.style_id for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH}

    def styled(self, paragraph: DocxParagraph, style: str, spans: layout.Spans = ()) -> DocxParagraph:
        paragraph._p.style = self.ids[style]
        return _runs(paragraph, spans, SIZE[style])

    def paragraph(self, style: str, spans: layout.Spans = ()) -> DocxParagraph:
        """A paragraph of the body; straight after a table it keeps a little space, and after lines left
        to write on, the space the PDF leaves under them as well as its own, as in the PDF."""
        paragraph = self.styled(self.doc.add_paragraph(), style, spans)
        if self.after_space:
            paragraph.paragraph_format.space_before = Pt(6 + BEFORE[style])
        elif self.after_table and BEFORE[style] < 6:
            paragraph.paragraph_format.space_before = Pt(6)
        self.after_table = self.after_space = False
        return paragraph

    def cell(self, cell: _Cell, spans: layout.Spans, *, style: str = "Table text", right: bool = False,
             fill: str | None = None) -> DocxParagraph:
        paragraph = cell.paragraphs[0]
        if right:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if fill:
            _shade(cell, fill)
        return self.styled(paragraph, style, spans)

    def block(self, b: layout.Block) -> list[list[DocxParagraph]]:
        if isinstance(b, layout.Title):
            return self.title(b)
        if isinstance(b, layout.Heading):
            heading = self.paragraph("Heading 1", (layout.Span(b.text),))
            if b.new_page:
                heading.paragraph_format.page_break_before = True
            return [[heading]]
        if isinstance(b, layout.Paragraph):
            paragraph = self.paragraph(BLOCK_STYLES[b.style], b.spans)
            if b.indent and b.style != "question":
                paragraph.paragraph_format.left_indent = INDENT
            if b.keep:
                paragraph.paragraph_format.keep_with_next = True
            return [[paragraph]]
        if isinstance(b, layout.Table):
            return self.table(b)
        if isinstance(b, layout.Space):
            return self.space(b)
        if isinstance(b, layout.Group):
            rows = [row for inner in b.blocks for row in self.block(inner)]
            _keep(rows[:-1])
            return rows
        return []

    def new_table(self, rows: int, cols: int) -> DocxTable:
        """Word joins two tables that touch: a table straight after another gets an empty line between."""
        if self.after_table:
            self.styled(self.doc.add_paragraph(), "Table text")
        self.after_table, self.after_space = True, False
        return self.doc.add_table(rows=rows, cols=cols)

    def title(self, t: layout.Title) -> list[list[DocxParagraph]]:
        made = [[self.paragraph("Memo title", (layout.Span(t.text),))]]
        if t.facts:
            table = self.new_table(len(t.facts), 2)
            _layout(table, (0.36, 0.64), (), padding=(20, 0, 30, 80))
            for row, (label, value) in zip(table.rows, t.facts):
                made.append([self.cell(row.cells[0], (layout.Span(label),), style="Title label"),
                             self.cell(row.cells[1], value, style="Title value")])
            self.after_table = False
        if t.note:
            note = self.paragraph("Memo note", (layout.Span(t.note, italic=True),))
            note.paragraph_format.space_before = Pt(5)
            made.append([note])
        rule = self.paragraph("Table text")
        rule.paragraph_format.space_after = Pt(6)
        _rule_below(rule)
        return made

    def table(self, t: layout.Table) -> list[list[DocxParagraph]]:
        made = []
        if t.title:
            made.append([self.paragraph("Table caption", (layout.Span(t.title),))])
        header = 1 if t.columns else 0
        table = self.new_table(header + len(t.rows), len(t.widths))
        _layout(table, t.widths, EDGES)
        if header:
            _repeat(table.rows[0])
            heads = [self.cell(cell, (layout.Span(text, bold=True),), right=j in t.numeric, fill=HEAD_FILL)
                     for j, (cell, text) in enumerate(zip(table.rows[0].cells, t.columns))]
            _keep([heads])  # the header row stays with the first row
            made.append(heads)
        rows = []
        for row, data in zip(table.rows[header:], t.rows):
            if t.row_height:
                _least(row, t.row_height)
            if layout.row_lines(data, t.widths) <= WHOLE_ROW:
                _whole(row)
            rows.append([self.cell(cell, spans, right=j in t.numeric,
                                   fill=LABEL_FILL if t.kind in ("facts", "form") and j == 0 else None)
                         for j, (cell, spans) in enumerate(zip(row.cells, data))])
        # A short table is kept on one page, as in the PDF; the lines under a table stay with its last row.
        if layout.table_height(t) <= t.whole:
            _keep(rows[:-1])
        under = []
        if t.note:
            self.after_table = False
            under.append([self.paragraph("Table note", (layout.Span(t.note, italic=True),))])
        for p in t.tail:
            under += self.block(p)
        if under:
            _keep(rows[-1:] + under[:-1])
        return made + rows + under

    def space(self, s: layout.Space) -> list[list[DocxParagraph]]:
        """Ruled lines LINE points apart, or a box as tall as so many lines, as in the PDF: the rows have
        no top or bottom padding, and what is written in them is set as body text."""
        table = self.new_table(1 if s.box else s.lines, 1)
        edges = ("top", "left", "bottom", "right") if s.box else ("bottom", "insideH")
        padding = (0, PADDING[1], 0, PADDING[3])
        _layout(table, (1.0,), edges, indent=INDENT if s.indent else 0, padding=padding)
        made = []
        for row in table.rows:
            _least(row, LINE * s.lines if s.box else LINE, padding)
            _whole(row)
            paragraph = self.styled(row.cells[0].paragraphs[0], "Normal")
            paragraph.paragraph_format.space_after = Pt(0)
            if s.box:
                paragraph.paragraph_format.space_before = Pt(3)
            made.append([paragraph])
        self.after_space = True
        return made


def _page(doc, document: layout.Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = section.bottom_margin = Cm(2)
    section.top_margin = Cm(2.3)
    section.header_distance = section.footer_distance = Cm(1)
    header = section.header.paragraphs[0]
    header.style = "Header"
    # On one line, as in the PDF: a smaller size first, then shortened.
    text, size = layout.fit(document.header, 8)
    run = header.add_run(text)
    if size != 8:
        run.font.size = Pt(size)
    _rule_below(header)
    footer = section.footer.paragraphs[0]
    footer.style = "Footer"
    footer.add_run(f"{document.footer} · page ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")


def _clip(text: str) -> str:
    """A property within the length Word allows."""
    return text if len(text) <= LIMIT else text[:LIMIT - 1] + "…"


def _properties(doc, document: layout.Document) -> None:
    """Title, subject and author, the configuration version in full, and the as_of date as the file's
    date: never the clock."""
    props = doc.core_properties
    props.title = _clip(document.title)
    props.subject = _clip(document.subject)
    props.author = props.last_modified_by = "CreditIQ"
    props.keywords = _clip(document.keywords)
    props.version = _clip(document.version)
    props.comments = props.category = ""
    props.revision = 1
    if document.as_of:
        props.created = props.modified = datetime(document.as_of.year, document.as_of.month, document.as_of.day)
    # The template's thumbnail is a picture of an empty page: the file goes without one.
    rels = doc.part.package.rels
    for rid in [rid for rid, rel in rels.items() if rel.reltype == RT.THUMBNAIL]:
        rels.pop(rid)


def _package(data: bytes, as_of: date | None) -> bytes:
    """The package again with every entry dated as_of, so the same case gives the same file, and the
    application properties CreditIQ's."""
    stamp = (as_of.year, as_of.month, as_of.day, 0, 0, 0) if as_of and as_of.year >= 1980 else (1980, 1, 1, 0, 0, 0)
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=stamp)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            target.writestr(entry, APP if info.filename == "docProps/app.xml" else source.read(info))
    return out.getvalue()


def render(document: layout.Document) -> bytes:
    doc = new_document()
    _styles(doc)
    _settings(doc)
    _page(doc, document)
    writer = _Writer(doc)
    for block in document.blocks:
        writer.block(block)
    _properties(doc, document)
    buffer = io.BytesIO()
    doc.save(buffer)
    return _package(buffer.getvalue(), document.as_of)
