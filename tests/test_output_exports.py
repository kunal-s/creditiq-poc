"""The draft CAM and the PD note as Word and PDF files (FRD F-23.5, F-24.3; TC-26, TC-27).

The writers are tested on a memo and a note built here, so that they do not depend on how a case is
drafted; the downloads end to end on a processed fixture case. Every name is fictional.
"""

import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document as read_docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from fastapi.testclient import TestClient
from PIL import ImageOps
from reportlab.pdfbase.pdfmetrics import stringWidth

from engine import auth, db, exports, pipeline
from engine.contracts import (
    Cam,
    CamBlock,
    CamFact,
    CamSection,
    CamSignOff,
    CamSummaryItem,
    CamTable,
    CamText,
    Evidence,
    OutputSource,
    PdNote,
    PdPerson,
    PdQuestion,
    PdRef,
    PdTopic,
)
from engine.exports import layout
from tests.conftest import ANALYST, MANAGER, RM, case_files, make_case, run_jobs, sign_in, upload

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ADMIN = "nikhil.sawant@rblbank.com"
SCREENS = "http://localhost:5173"
COI, AFS, GST, ITR, BANK, MSG, KYC = "D-COI", "D-AFS", "D-GST", "D-ITR", "D-BANK", "D-MSG", "D-KYC"
HASH = "4ba5d6b3731946bdab564fdf526c87cc7ef97bb2054d34878b3420190474b424"
LABEL = "Version 7 (published 29 Sep 2026)"
PROMPTS = ["Recommendation", "Terms and conditions", "Deviations recommended for approval",
           "Decision of the approving authority"]
ROLES = ["Prepared by", "Recommended by", "Approved by"]
RECORD = ["Date of discussion", "Place", "Bank officials present", "Promoters or representatives present"]
TOP, FOOT = 2.3 / 2.54 * 72, 2 / 2.54 * 72
"""Points above and below the text of a page: the margins, the running header and the footer outside them."""


def _t(text: str, *cited: tuple[str, int]) -> CamText:
    return CamText(text=text, citations=[Evidence(document_id=d, page=p) for d, p in cited])


def _source(number: int, document_id: str, label: str, file_name: str, pages: list[int]) -> OutputSource:
    return OutputSource(number=number, document_id=document_id, label=label, file_name=file_name, pages=pages)


def _memo(**changes) -> Cam:
    sections = [
        CamSection(id="snapshot", number=1, title="Proposal at a glance", blocks=[CamBlock(kind="facts", facts=[
            CamFact(label="Borrower", value=_t("Kestrel Fabricators Pvt Ltd", (COI, 1))),
            CamFact(label="Amount requested", value=_t("INR 4.00 cr")),
        ])]),
        CamSection(id="financials", number=2, title="Financial performance", blocks=[
            CamBlock(kind="table", table=CamTable(
                columns=["", "FY 2024-25 (audited)", "FY 2025-26 (audited)"], numeric=[1, 2],
                note="Figures as on the spread (Outputs → Spread).", rows=[
                    [_t("Revenue from operations"), _t("INR 41.06 cr", (AFS, 4)),
                     _t("INR 45.92 cr", (AFS, 4), (AFS, 5))],
                    [_t("Current ratio"), _t("1.28x", (AFS, 3)), _t("1.19x", (AFS, 3))],
                ])),
            CamBlock(kind="paragraph", paragraph=_t("Revenue from operations grew 11.8%, against GST outward supplies "
                                                    "3.4% lower.", (AFS, 4), (GST, 2))),
            CamBlock(kind="subheading", text="Liquidity"),
            CamBlock(kind="note", text="The current ratio is below the norm of at least 1.20x: see section 3."),
        ]),
        CamSection(id="banking", number=3, title="Banking and account conduct", empty="Nothing on file yet."),
    ]
    summary = [
        CamSummaryItem(kind="finding", ref="XV-01", rule="ID-01", title="PAN differs between documents",
                       severity="serious", status="Open", blocking=True,
                       text="The income tax return shows PAN AAACK9876Q against AAACK1234F on the company PAN.",
                       citations=[Evidence(document_id=ITR, page=1)]),
        CamSummaryItem(kind="incomplete", ref="XV-02", rule="ID-04", title="Promoters match the KYC received",
                       text="Director KYC not on file for Sunil Karve.", request="the director KYC of Sunil Karve"),
        CamSummaryItem(kind="deviation", ref="FR-01", rule="FR-01", title="Current ratio",
                       text="Category B: Current ratio 1.19x against at least 1.20x.",
                       citations=[Evidence(document_id=AFS, page=3)]),
        CamSummaryItem(kind="unassessed", ref="DOC-02", rule="DOC-02", title="Statements for 12 months",
                       text="Statements for 12 months could not be assessed: GSTR-3B for Dec 2025 not on file."),
        CamSummaryItem(kind="query", ref="Q-1", group="documents_needed", text="GSTR-3B for Dec 2025 to Mar 2026."),
        CamSummaryItem(kind="query", ref="Q-2", group="documents_to_redo",
                       text="Bank statement: outdated. Please provide one dated within the last 60 days.",
                       citations=[Evidence(document_id=BANK, page=2)]),
        CamSummaryItem(kind="query", ref="Q-3", group="clarifications",
                       text="Please explain the credit of INR 50.00 lakh on 20 Jun 2025.",
                       citations=[Evidence(document_id=BANK, page=2)]),
    ]
    sources = [
        _source(1, COI, "Certificate of incorporation", "kestrel certificate of incorporation.pdf", [1]),
        _source(2, AFS, "Audited financial statements · FY 2025-26", "kestrel audited fy2026.pdf", [3, 4, 5]),
        _source(3, GST, "GSTR-3B · Sep 2025", "kestrel gst returns.pdf", [2]),
        _source(4, ITR, "Income tax return", "kestrel itr.pdf", [1]),
        _source(5, BANK, "Bank statement · Lotuscrest Bank", "kestrel current account.pdf", [2]),
    ]
    memo = Cam(case_id="BBG-2026-000004", title="Credit Approval Memo", borrower="Kestrel Fabricators Pvt Ltd",
               as_of="2026-09-29", config_version=HASH, config_label=LABEL, run_id="XR-0001",
               interim_template=True, template_note="Draft — layout assumed until RBL's CAM template is received.",
               sections=sections, summary_number=4, summary_title="Summary of findings, deviations and open queries",
               summary=summary, recommendation_number=5, recommendation_title="Recommendation",
               recommendation_prompts=PROMPTS, sign_off=[CamSignOff(role=r) for r in ROLES], sources=sources)
    return memo.model_copy(update=changes)


def _question(number: int, area: str, question: str, context: str, **fields) -> PdQuestion:
    return PdQuestion(id=f"PDQ-{number}", number=number, question=question, context=context,
                      text=f"{context} {question}", area=area, **fields)


def _note(**changes) -> PdNote:
    questions = [
        _question(1, "identity", "Why does the income tax return show a different PAN?",
                  "The income tax return shows PAN AAACK9876Q against AAACK1234F on the company PAN.",
                  severity="serious", source="finding", ref="F-1", xv_ref="XV-01", rule="ID-01",
                  title="PAN differs between documents", status="Open", target="cross_verification",
                  evidence=[Evidence(document_id=ITR, page=1), Evidence(document_id=COI, page=2)]),
        _question(2, "turnover", "Where are the rest of the sales receipts banked, or are they received in cash?",
                  "GST outward supplies run 25.0% above the business credits in the bank (XV-04); the income tax "
                  "return runs 24.4% above them (XV-05).",
                  severity="moderate", source="finding", ref="F-4", xv_ref="XV-04", rule="TO-02",
                  title="GST turnover not supported by bank credits", status="Accepted as valid",
                  target="cross_verification", evidence=[Evidence(document_id=GST, page=2)],
                  also=[PdRef(ref="F-5", xv_ref="XV-05", rule="TX-05", status="Open",
                              title="Turnover in the tax return not supported by bank credits",
                              evidence=[Evidence(document_id=BANK, page=2)])]),
        # An observation: a fact the memo states that needs the promoter's answer, asked after the
        # findings and deviations of its severity.
        _question(3, "obligations", "Which lender disbursed the INR 50.00 lakh on 20 Jun 2025, on what terms, and "
                                    "why is no instalment seen?",
                  "A credit of INR 50.00 lakh on 20 Jun 2025 (TL DISBURSAL LN 778812) matches no facility on file.",
                  severity="moderate", source="observation", ref="obligations:loan-disbursal", rule="OBS",
                  title="Loan disbursal no facility explains", status="section 9", target="cam",
                  evidence=[Evidence(document_id=BANK, page=2)]),
        _question(4, "policy", "What is the plan to bring the current ratio up to the norm?",
                  "Current ratio is 1.19x in FY 2025-26 against a norm of at least 1.20x.",
                  severity="moderate", source="deviation", ref="FR-01", rule="FR-01", title="Current ratio",
                  target="policy", evidence=[Evidence(document_id=AFS, page=3)]),
    ]
    note = PdNote(
        case_id="BBG-2026-000004", title="Personal discussion — question note",
        borrower="Kestrel Fabricators Pvt Ltd", as_of="2026-09-29", config_version=HASH, config_label=LABEL,
        header=[CamFact(label="Borrower", value=_t("Kestrel Fabricators Pvt Ltd")),
                CamFact(label="Case ID", value=_t("BBG-2026-000004")),
                CamFact(label="Constitution", value=_t("Private limited")),
                CamFact(label="Proposal", value=_t("Cash credit and term loan, INR 6.50 cr", (MSG, 1))),
                CamFact(label="Data as of", value=_t("29 Sep 2026")),
                CamFact(label="Configuration version", value=_t(LABEL))],
        people=[PdPerson(name="Tarun Velankar", role="Director", din="07812345",
                         sources=["Certificate of incorporation", "Director KYC"],
                         evidence=[Evidence(document_id=COI, page=1), Evidence(document_id=KYC, page=1)]),
                PdPerson(name="Sunil Karve", role="Guarantor", sources=["Sourcing message"],
                         evidence=[Evidence(document_id=MSG, page=1)])],
        questions=questions,
        topics=[PdTopic(area="identity", title="Identity and KYC", question_ids=["PDQ-1"]),
                PdTopic(area="turnover", title="Turnover, tax and sales", question_ids=["PDQ-2"]),
                PdTopic(area="obligations", title="Existing borrowings", question_ids=["PDQ-3"]),
                PdTopic(area="policy", title="Policy norms", question_ids=["PDQ-4"])],
        documents_to_collect=[_t("KYC of Sunil Karve."), _t("Bank statement: outdated.", (BANK, 2))],
        record_fields=RECORD, sign_off=[CamSignOff(role="Discussion held by")],
        # Numbered in reading order: the header, the people, the questions (and what they also cover).
        sources=[_source(1, MSG, "Sourcing message", "application-message.txt", [1]),
                 _source(2, COI, "Certificate of incorporation", "kestrel certificate of incorporation.pdf", [1, 2]),
                 _source(3, KYC, "Director KYC · Tarun Velankar", "kestrel kyc tarun.pdf", [1]),
                 _source(4, ITR, "Income tax return", "kestrel itr.pdf", [1]),
                 _source(5, GST, "GSTR-3B · Sep 2025", "kestrel gst returns.pdf", [2]),
                 _source(6, BANK, "Bank statement · Lotuscrest Bank", "kestrel current account.pdf", [2]),
                 _source(7, AFS, "Audited financial statements · FY 2025-26", "kestrel audited fy2026.pdf", [3])],
    )
    return note.model_copy(update=changes)


# --- Reading the files back ---

def _n(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def _squash(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _body(data: bytes) -> list[tuple]:
    """The Word file's body in order: ("p", style, text) for a paragraph, ("t", rows of cell texts) for a table."""
    out = []
    for item in read_docx(io.BytesIO(data)).iter_inner_content():
        if isinstance(item, DocxTable):
            out.append(("t", [[_n(c.text) for c in row.cells] for row in item.rows]))
        else:
            out.append(("p", item.style.name, _n(item.text)))
    return out


def _texts(body: list[tuple]) -> list[str]:
    return [t for b in body for t in ([b[2]] if b[0] == "p" else [c for row in b[1] for c in row]) if t]


def _headings(body: list[tuple]) -> list[str]:
    return [b[2] for b in body if b[0] == "p" and b[1] == "Heading 1"]


def _pages(data: bytes) -> list[str]:
    pdf = pdfium.PdfDocument(data)
    try:
        return [page.get_textpage().get_text_range() for page in pdf]
    finally:
        pdf.close()


def _running(line: str) -> str:
    """A running line as the PDF draws it: fitted to the width, in the standard fonts' characters."""
    return _squash(layout.winansi(layout.fit(line, 8)[0]))


def _pdf_body(data: bytes, document: exports.Document) -> str:
    """The PDF's text without whitespace, the running header and footer taken off every page."""
    pages = _pages(data)
    out = []
    for n, text in enumerate(pages, start=1):
        text = _squash(text)
        for line in (document.header, f"{document.footer} · page {n} of {len(pages)}"):
            assert _running(line) in text, (n, line)
            text = text.replace(_running(line), "", 1)
        out.append(text)
    return "".join(out)


def _word_as_pdf(data: bytes) -> str:
    """The Word file's text as the PDF writes it: the same text, in the standard fonts' characters."""
    return _squash(layout.winansi("".join(_texts(_body(data)))))


def _after(body: list[tuple], text: str) -> tuple:
    index = next(i for i, b in enumerate(body) if b[0] == "p" and b[2] == text)
    return body[index + 1]


def _page_of(pages: list[str], text: str) -> int:
    return next(i for i, page in enumerate(pages) if text in page)


def _row(page: str, label: str) -> bool:
    """Whether the page holds the table row of this label ("Row 1", not "Row 10")."""
    return re.search(rf"\b{label}\b", page) is not None


# --- The writers ---

def test_the_word_memo_follows_the_memo():
    memo = _memo()
    body = _body(exports.render(exports.cam_document(memo), "docx"))
    assert body[0] == ("p", "Memo title", "Credit Approval Memo")
    # The configuration by the name people know it by, not its hash.
    assert body[1][1] == [["Borrower", "Kestrel Fabricators Pvt Ltd"], ["Case ID", "BBG-2026-000004"],
                          ["Drafted from the data on file as of", "29 Sep 2026"], ["Configuration version", LABEL],
                          ["Cross-check run", "XR-0001"]]
    assert ("p", "Memo note", "Draft — layout assumed until RBL's CAM template is received.") in body
    assert _headings(body) == ["1. Proposal at a glance", "2. Financial performance", "3. Banking and account conduct",
                               "4. Summary of findings, deviations and open queries", "5. Recommendation", "Sources"]
    # Facts as label and value, the citation after the value; a table's figures cited cell by cell.
    assert _after(body, "1. Proposal at a glance")[1] == [["Borrower", "Kestrel Fabricators Pvt Ltd [1 p.1]"],
                                                         ["Amount requested", "INR 4.00 cr"]]
    assert _after(body, "2. Financial performance")[1] == [
        ["", "FY 2024-25 (audited)", "FY 2025-26 (audited)"],
        ["Revenue from operations", "INR 41.06 cr [2 p.4]", "INR 45.92 cr [2 pp.4-5]"],
        ["Current ratio", "1.28x [2 p.3]", "1.19x [2 p.3]"]]
    # The Word file keeps the characters as written (the PDF writes the arrow as ->).
    assert ("p", "Table note", "Figures as on the spread (Outputs → Spread).") in body
    assert ("p", "Normal", "Revenue from operations grew 11.8%, against GST outward supplies 3.4% lower. "
                           "[2 p.4; 3 p.2]") in body
    assert ("p", "Memo subheading", "Liquidity") in body
    assert _after(body, "3. Banking and account conduct") == ("p", "Memo note", "Nothing on file yet.")


def test_the_summary_lists_findings_as_a_register_then_checks_not_completed_deviations_norms_and_queries():
    memo = _memo()
    body = _body(exports.render(exports.cam_document(memo), "docx"))
    subheadings = [b[2] for b in body if b[0] == "p" and b[1] == "Memo subheading"]
    assert subheadings[-5:] == ["Findings", "Checks not completed", "Deviations from policy norms",
                                "Norms not assessed", "Open queries"]
    # A finding's explanation is said once, in the cross-verification section: the summary is a register.
    assert _after(body, "Findings")[1] == [
        ["Ref", "Severity", "Finding", "Status"],
        ["XV-01 ID-01", "Serious", "PAN differs between documents [4 p.1]", "Open Blocks submission"]]
    assert not any("AAACK9876Q" in t for t in _texts(body))
    assert _after(body, "Checks not completed")[1] == [
        ["Ref", "Check", "What would complete it"],
        ["XV-02 ID-04", "Promoters match the KYC received", "the director KYC of Sunil Karve"]]
    assert _after(body, "Deviations from policy norms")[1][1] == [
        "FR-01", "Current ratio Category B: Current ratio 1.19x against at least 1.20x. [2 p.3]"]
    assert _after(body, "Norms not assessed")[1][0] == ["Ref", "Norm"]
    # Open queries under the groups of the query list, numbered across them.
    assert _after(body, "Documents needed")[1] == [["No.", "Query"], ["1", "GSTR-3B for Dec 2025 to Mar 2026."]]
    assert _after(body, "Documents to redo")[1][1] == [
        "2", "Bank statement: outdated. Please provide one dated within the last 60 days. [5 p.2]"]
    assert _after(body, "Clarifications")[1][1] == ["3", "Please explain the credit of INR 50.00 lakh on 20 Jun 2025. "
                                                         "[5 p.2]"]


def test_with_no_finding_the_summary_says_so_for_the_run():
    memo = _memo(summary=[i for i in _memo().summary if i.kind != "finding"])
    body = _body(exports.render(exports.cam_document(memo), "docx"))
    assert _after(body, "Findings") == ("p", "Memo note", "No discrepancy in cross-check run XR-0001.")
    nothing = _body(exports.render(exports.cam_document(_memo(run_id=None, summary=[])), "docx"))
    assert "Findings" not in _texts(nothing)
    assert ("p", "Memo note", "No findings, deviations or open queries.") in nothing


def test_the_memo_ends_in_the_blank_recommendation_the_sign_off_and_the_sources():
    memo = _memo()
    body = _body(exports.render(exports.cam_document(memo), "docx"))
    # Each prompt with an empty box under it, then the sign-off with every cell blank.
    for prompt in PROMPTS:
        box = _after(body, prompt)
        assert box[0] == "t" and box[1] == [[""]], prompt
    sign_off = next(b[1] for b in body if b[0] == "t" and b[1][0] == ["", "Name", "Designation", "Signature", "Date"])
    assert [row[0] for row in sign_off[1:]] == ROLES and all(c == "" for row in sign_off[1:] for c in row[1:])
    sources = _after(body, "Sources")[1]
    assert sources[0] == ["No.", "Document", "File", "Pages"]
    assert [int(row[0]) for row in sources[1:]] == [s.number for s in memo.sources]
    assert sources[2] == ["2", "Audited financial statements · FY 2025-26", "kestrel audited fy2026.pdf", "3-5"]


def test_the_word_note_asks_each_question_with_room_for_the_answer():
    note = _note()
    body = _body(exports.render(exports.pd_note_document(note), "docx"))
    assert body[0] == ("p", "Memo title", "Personal discussion — question note")
    # The header as the note gives it, cited where a document backs it.
    assert body[1][1] == [["Borrower", "Kestrel Fabricators Pvt Ltd"], ["Case ID", "BBG-2026-000004"],
                          ["Constitution", "Private limited"], ["Proposal", "Cash credit and term loan, INR 6.50 cr [1 p.1]"],
                          ["Data as of", "29 Sep 2026"], ["Configuration version", LABEL]]
    assert _headings(body) == ["People to meet", "Record of the discussion", "Identity and KYC",
                               "Turnover, tax and sales", "Existing borrowings", "Policy norms",
                               "Documents to collect at the discussion", "Sources"]
    # Each person with the pages that name them.
    assert _after(body, "People to meet")[1] == [
        ["Name", "Role", "DIN", "Named in"],
        ["Tarun Velankar", "Director", "07812345", "Certificate of incorporation, Director KYC [2 p.1; 3 p.1]"],
        ["Sunil Karve", "Guarantor", "—", "Sourcing message [1 p.1]"]]
    assert _after(body, "Record of the discussion")[1] == [[f, ""] for f in RECORD]
    questions = [b[2] for b in body if b[0] == "p" and b[1] == "Memo question"]
    assert questions == [f"{q.number}. {q.question}" for q in note.questions]
    assert ("p", "Memo detail", "Serious · Finding XV-01 (ID-01) · Open") in body
    # A question asked for several findings names the others it covers, and cites their pages too.
    assert ("p", "Memo detail", "Moderate · Finding XV-04 (TO-02) · Accepted as valid "
                                "Also covers: XV-05 (TX-05) · Open") in body
    assert ("p", "Normal", "Evidence: [5 p.2; 6 p.2]") in body
    assert ("p", "Memo detail", "Moderate · Deviation from FR-01") in body
    assert ("p", "Normal", f"Why we ask: {note.questions[1].context}") in body
    assert ("p", "Normal", "Evidence: [2 p.2; 4 p.1]") in body
    # Four blank lines for the response and two for the follow-up under every question.
    responses = [i for i, b in enumerate(body) if b[0] == "p" and b[2] == "Response"]
    follow_ups = [i for i, b in enumerate(body) if b[0] == "p" and b[2] == "Follow-up"]
    assert len(responses) == len(follow_ups) == len(note.questions)
    assert all(body[i + 1] == ("t", [[""]] * 4) for i in responses)
    assert all(body[i + 1] == ("t", [[""]] * 2) for i in follow_ups)
    assert _after(body, "Documents to collect at the discussion")[1] == [
        ["No.", "Document"], ["1", "KYC of Sunil Karve."], ["2", "Bank statement: outdated. [6 p.2]"]]
    sign_off = next(b[1] for b in body if b[0] == "t" and b[1][0][:2] == ["", "Name"])
    assert sign_off[1:] == [["Discussion held by", "", "", "", ""]]
    assert [int(r[0]) for r in _after(body, "Sources")[1][1:]] == [s.number for s in note.sources]


def test_an_observation_is_asked_by_its_title_as_a_fact_the_memo_states():
    note = _note()
    observation = note.questions[2]
    body = _body(exports.render(exports.pd_note_document(note), "docx"))
    # Under its topic, with the memo section that states it (its status): its key and its rule (OBS)
    # mean nothing to the reader.
    meta = "Moderate · Observation · section 9"
    assert _after(body, "Existing borrowings") == ("p", "Memo question", f"3. {observation.question}")
    assert _after(body, f"3. {observation.question}") == ("p", "Memo detail", meta)
    assert _after(body, meta) == ("p", "Normal", f"Why we ask: {observation.context}")
    assert _after(body, f"Why we ask: {observation.context}") == ("p", "Normal", "Evidence: [6 p.2]")
    assert not any("OBS" in t or observation.ref in t for t in _texts(body))
    # The word from the note's labels where they give one; no section where the memo numbers none.
    changed = note.model_copy(update={"labels": {"observation": "Memo observation"}, "questions": [
        *note.questions[:2], observation.model_copy(update={"status": None}), note.questions[3]]})
    body = _body(exports.render(exports.pd_note_document(changed), "docx"))
    assert ("p", "Memo detail", "Moderate · Memo observation") in body


def test_a_note_without_header_facts_or_sign_off_still_names_the_case():
    body = _body(exports.render(exports.pd_note_document(_note(header=[], sign_off=[], config_label="")), "docx"))
    assert body[1][1] == [["Borrower", "Kestrel Fabricators Pvt Ltd"], ["Case ID", "BBG-2026-000004"],
                          ["Data as of", "29 Sep 2026"], ["Configuration version", HASH[:8]]]
    assert not any(b[0] == "t" and b[1][0][:2] == ["", "Name"] for b in body)


def test_the_words_inside_the_files_come_from_the_labels_with_todays_wording_by_default():
    memo = _memo(labels={"case_id": "Application", "summary_findings": "Discrepancies found",
                         "summary_incomplete": "Checks to complete", "sources": "Documents cited", "source_no": "#",
                         "sign_name": "Full name"})
    body = _body(exports.render(exports.cam_document(memo), "docx"))
    assert ["Application", "BBG-2026-000004"] in body[1][1] and ["Configuration version", LABEL] in body[1][1]
    assert ("p", "Memo subheading", "Discrepancies found") in body and ("p", "Memo subheading", "Checks to complete") in body
    assert _headings(body)[-1] == "Documents cited"
    assert _after(body, "Documents cited")[1][0] == ["#", "Document", "File", "Pages"]
    assert any(b[0] == "t" and b[1][0] == ["", "Full name", "Designation", "Signature", "Date"] for b in body)
    # A label written with or without its colon leads into its text the same way.
    note = _note(labels={"people_to_meet": "Promoters to meet", "why_we_ask": "Reason:", "evidence": "Pages",
                         "also_covers": "Covers too", "response": "Answer", "follow_up": "Next steps",
                         "documents_to_collect": "To collect"})
    body = _body(exports.render(exports.pd_note_document(note), "docx"))
    assert _headings(body)[:2] == ["Promoters to meet", "Record of the discussion"] and "To collect" in _headings(body)
    assert ("p", "Normal", f"Reason: {note.questions[0].context}") in body
    assert ("p", "Normal", "Pages: [2 p.2; 4 p.1]") in body
    assert ("p", "Memo detail", "Moderate · Finding XV-04 (TO-02) · Accepted as valid Covers too: XV-05 (TX-05) · Open") in body
    assert sum(1 for b in body if b[0] == "p" and b[2] in ("Answer", "Next steps")) == 2 * len(note.questions)


@pytest.mark.parametrize("output", ["cam", "pd_note"])
def test_both_files_carry_the_same_text(output: str):
    document = exports.cam_document(_memo()) if output == "cam" else exports.pd_note_document(_note())
    word = exports.render(document, "docx")
    pdf = exports.render(document, "pdf")
    assert _pdf_body(pdf, document) == _word_as_pdf(word)
    pages = _pages(pdf)
    assert "page 1 of" in pages[0] and f"page {len(pages)} of {len(pages)}" in pages[-1]


def test_the_markers_number_the_sources_and_a_document_not_listed_is_added():
    memo = _memo(sources=_memo().sources[:2])  # GST, ITR and BANK are cited but not listed
    document = exports.cam_document(memo)
    body = _body(exports.render(document, "docx"))
    assert ("p", "Normal", "Revenue from operations grew 11.8%, against GST outward supplies 3.4% lower. "
                           "[2 p.4; 3 p.2]") in body
    assert [r[:2] + r[3:] for r in _after(body, "Sources")[1][1:]] == [
        ["1", "Certificate of incorporation", "1"], ["2", "Audited financial statements · FY 2025-26", "3-5"],
        ["3", "Document", "2"], ["4", "Document", "1"], ["5", "Document", "2"]]
    no_sources = _body(exports.render(exports.cam_document(_memo(sources=[], sections=[], summary=[])), "docx"))
    assert _after(no_sources, "Sources") == ("p", "Memo note", "No document is cited.")


def test_three_or_more_sources_in_a_row_at_the_same_pages_read_as_a_range():
    returns = [_source(n, f"D-GST-{n}", "GSTR-3B", f"gstr-3b-{n}.pdf", [1]) for n in range(1, 8)]
    statement = _t("GST outward supplies INR 12.00 cr.", *[(f"D-GST-{n}", 1) for n in (7, 1, 2, 3, 4, 6)],
                   ("D-GST-5", 2))
    memo = _memo(sources=returns, summary=[], sections=[CamSection(
        id="turnover", number=1, title="Turnover across sources",
        blocks=[CamBlock(kind="paragraph", paragraph=statement),
                CamBlock(kind="paragraph", paragraph=_t("Two returns.", ("D-GST-1", 1), ("D-GST-2", 1)))])])
    texts = _texts(_body(exports.render(exports.cam_document(memo), "docx")))
    assert "GST outward supplies INR 12.00 cr. [1-4 p.1; 5 p.2; 6 p.1; 7 p.1]" in texts
    assert "Two returns. [1 p.1; 2 p.1]" in texts


def test_the_word_file_keeps_every_character_and_the_pdf_writes_them_in_its_fonts():
    name = "Ğülsüm Ārav देवेंद्र Velankar"
    text = f"{name}: ratio ≥ 1.20x, ≤ 3.00x, ≠ 0 → met ✓; drop − 3%; ₹5,000; R&D <b>capex</b>\x07\x0b"
    kept = f"{name}: ratio ≥ 1.20x, ≤ 3.00x, ≠ 0 → met ✓; drop − 3%; ₹5,000; R&D <b>capex</b>"
    assert layout.clean(text) == kept  # control characters, which a Word file cannot hold, are left out
    assert layout.winansi(kept) == ("Gülsüm Arav ???????? Velankar: ratio >= 1.20x, <= 3.00x, != 0 -> met Yes; "
                                    "drop - 3%; INR 5,000; R&D <b>capex</b>")
    assert layout.winansi("≮") == "?" and layout.winansi("a​b­c") == "abc"
    memo = _memo(borrower=f"{name} Exports Pvt Ltd", sections=[CamSection(
        id="snapshot", number=1, title="Proposal at a glance → ✓",
        blocks=[CamBlock(kind="paragraph", paragraph=_t(text, (COI, 1)))])])
    document = exports.cam_document(memo)
    data = exports.render(document, "docx")
    words = _texts(_body(data))
    assert f"{kept} [1 p.1]" in words and "1. Proposal at a glance → ✓" in words
    doc = read_docx(io.BytesIO(data))
    assert doc.sections[0].header.paragraphs[0].text == document.header
    assert document.header == f"Credit Approval Memo — Draft · {name} Exports Pvt Ltd · BBG-2026-000004"
    assert doc.core_properties.subject == f"{name} Exports Pvt Ltd · BBG-2026-000004"
    # The PDF draws the same text in the standard fonts' characters; the markup comes through as text.
    pdf = exports.render(document, "pdf")
    assert _pdf_body(pdf, document) == _word_as_pdf(data)
    assert "Gülsüm Arav ???????? Velankar: ratio >= 1.20x" in _n(_pages(pdf)[0])


def test_the_configuration_is_named_by_its_label_and_the_files_carry_its_version_in_full():
    document = exports.cam_document(_memo())
    data = exports.render(document, "docx")
    props = read_docx(io.BytesIO(data)).core_properties
    assert (props.version, props.keywords) == (HASH, f"Configuration version {HASH}")
    assert not any(HASH in t for t in _texts(_body(data)))
    pdf = exports.render(document, "pdf")
    doc = pdfium.PdfDocument(pdf)
    try:
        assert doc.get_metadata_dict()["Keywords"] == f"Configuration version {HASH}"
    finally:
        doc.close()
    assert not any(HASH in _squash(page) for page in _pages(pdf))
    # Without a label, the start of the hash.
    body = _body(exports.render(exports.cam_document(_memo(config_label="")), "docx"))
    assert ["Configuration version", HASH[:8]] in body[1][1]


def test_the_files_are_laid_out_for_a4_with_a_running_header_and_page_numbers():
    document = exports.cam_document(_memo())
    data = exports.render(document, "docx")
    doc = read_docx(io.BytesIO(data))
    section = doc.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)
    assert round(section.left_margin.cm, 1) == round(section.right_margin.cm, 1) == 2.0
    assert section.header.paragraphs[0].text == document.header
    assert document.header == "Credit Approval Memo — Draft · Kestrel Fabricators Pvt Ltd · BBG-2026-000004"
    footer = section.footer.paragraphs[0]._p
    assert [i.text.strip() for i in footer.iter(qn("w:instrText"))] == ["PAGE", "NUMPAGES"]
    assert "CreditIQ draft · data as of 29 Sep 2026 · page" in section.footer.paragraphs[0].text
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "" and len(t.columns) == 3)
    assert table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert table.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "E6EBF1"
    # Word's current layout rules, widow and orphan control, and CreditIQ as the application.
    compat = {s.get(qn("w:name")): s.get(qn("w:val")) for s in doc.settings.element.iter(qn("w:compatSetting"))}
    assert compat["compatibilityMode"] == "15"
    assert doc.styles["Normal"].paragraph_format.widow_control is True
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert "<Application>CreditIQ</Application>" in package.read("docProps/app.xml").decode()
        assert not [n for n in package.namelist() if "thumbnail" in n]
    pdf = pdfium.PdfDocument(exports.render(document, "pdf"))
    try:
        assert [tuple(round(x) for x in page.get_size()) for page in pdf] == [(595, 842)] * len(pdf)
    finally:
        pdf.close()


def test_word_rules_its_lines_and_sizes_its_rows_as_the_pdf_does():
    doc = read_docx(io.BytesIO(exports.render(exports.pd_note_document(_note()), "docx")))

    def padding(table: DocxTable) -> tuple[str, str]:
        margins = table._tbl.tblPr.find(qn("w:tblCellMar"))
        return margins.find(qn("w:top")).get(qn("w:w")), margins.find(qn("w:bottom")).get(qn("w:w"))

    # Word adds a row's top and bottom padding and its border to the height it is given: the ruled lines
    # have none, and 19.5 pt plus the border puts them 20 pt apart, as in the PDF. What is written on
    # them is body text.
    lines = [t for t in doc.tables if len(t.columns) == 1]
    assert len(lines) == 2 * len(_note().questions)
    for t in lines:
        assert padding(t) == ("0", "0")
        assert [row.height.pt for row in t.rows] == [19.5] * len(t.rows)
        assert {row.cells[0].paragraphs[0].style.name for row in t.rows} == {"Normal"}
    # The sign-off rows 30 pt tall and the record's 24 pt, the padding and border taken off.
    sign_off = next(t for t in doc.tables if len(t.columns) == 5 and t.rows[0].cells[1].text == "Name")
    assert padding(sign_off) == ("50", "70") and sign_off.rows[1].height.pt == 23.5
    record = next(t for t in doc.tables if len(t.columns) == 2 and t.rows[0].cells[0].text == RECORD[0])
    assert {row.height.pt for row in record.rows} == {17.5}
    # The memo's boxes as tall as five lines, for credit to write in as body text.
    doc = read_docx(io.BytesIO(exports.render(exports.cam_document(_memo()), "docx")))
    boxes = [t for t in doc.tables if len(t.columns) == 1]
    assert len(boxes) == len(PROMPTS)
    assert all(t.rows[0].height.pt == 99.5 and t.rows[0].cells[0].paragraphs[0].style.name == "Normal" for t in boxes)


def test_word_keeps_a_short_table_on_one_page_and_the_line_after_a_table_with_its_last_row():
    doc = read_docx(io.BytesIO(exports.render(exports.pd_note_document(_note()), "docx")))
    people = next(t for t in doc.tables if t.rows[0].cells[0].text == "Name")

    def kept(row) -> set[bool]:
        return {bool(p.paragraph_format.keep_with_next) for cell in row.cells for p in cell.paragraphs}

    assert [kept(row) for row in people.rows] == [{True}, {True}, {False}]
    # The financial table runs into its note and the line after it: its last row stays with them.
    doc = read_docx(io.BytesIO(exports.render(exports.cam_document(_memo()), "docx")))
    figures = next(t for t in doc.tables if t.rows[0].cells[0].text == "" and len(t.columns) == 3)
    assert kept(figures.rows[-1]) == {True}
    note = next(p for p in doc.paragraphs if p.text.startswith("Figures as on the spread"))
    after = next(p for p in doc.paragraphs if p.text.startswith("Revenue from operations grew"))
    assert note.paragraph_format.keep_with_next and not after.paragraph_format.keep_with_next


def test_word_does_not_break_a_line_inside_a_reference():
    memo = _memo(sources=[_source(n, d, "Document", f"{d}.pdf", [1]) for n, d in enumerate((COI, AFS, GST), start=1)],
                 summary=[], sections=[CamSection(id="turnover", number=1, title="Turnover", blocks=[
                     CamBlock(kind="paragraph", paragraph=_t(
                         "Finding XV-000004-12 on GSTR-3B for FY 2025-26, from a cross-check follow-up.",
                         (COI, 1), (AFS, 1), (GST, 1)))])])
    doc = read_docx(io.BytesIO(exports.render(exports.cam_document(memo), "docx")))
    paragraph = next(p for p in doc.paragraphs if p.text.startswith("Finding XV-"))
    assert _n(paragraph.text) == ("Finding XV-000004-12 on GSTR-3B for FY 2025-26, from a cross-check follow-up. "
                                  "[1-3 p.1]")
    # XV-000004-12 (two), GSTR-3B, 2025-26 and the range of sources 1-3: hyphens that do not break.
    assert len(paragraph._p.findall(f".//{qn('w:noBreakHyphen')}")) == 5
    assert "".join(t.text for t in paragraph._p.iter(qn("w:t"))).count("-") == 2  # cross-check, follow-up


def test_figures_and_short_lines_keep_to_one_line_and_the_columns_that_wrap_give_up_the_room():
    banking = CamTable(columns=["Account", "Months", "Gross credits", "Excluded", "Business credits",
                                "Average a month"], numeric=[2, 3, 4, 5], rows=[
        [_t("Lotuscrest Bank ·1234"), _t("12 (Apr 2025 to Mar 2026)"), _t("INR 7.72 cr", (BANK, 1)),
         _t("INR 52.00 lakh", (BANK, 1)), _t("INR 7.20 cr", (BANK, 1)), _t("INR 60.00 lakh", (BANK, 1))],
        [_t("Ironleaf Bank ·7777"), _t("12 (Apr 2025 to Mar 2026)"), _t("INR 1.80 cr", (GST, 1)), _t("INR 0"),
         _t("INR 1.80 cr", (GST, 1)), _t("INR 15.00 lakh", (GST, 1))]])
    facilities = CamTable(columns=["Lender", "Facility", "Declared", "Sanction letter", "Bureau", "Instalment",
                                   "Seen in the bank (a month)"], numeric=[2, 3, 4, 5, 6], rows=[
        [_t("Ashwood Finserv Limited"), _t("Term loan"), _t("INR 80.00 lakh", (AFS, 1)), _t("INR 95.00 lakh", (ITR, 1)),
         _t("INR 95.00 lakh; outstanding INR 60.00 lakh", (COI, 1)), _t("INR 1.45 lakh", (AFS, 1)),
         _t("INR 1.45 lakh (12 months)", (BANK, 1))],
        [_t("Tidewater Finance Limited"), _t("Term loan"), _t("Not declared"), _t("—"),
         _t("INR 1.20 cr; outstanding INR 75.00 lakh", (COI, 1)), _t("—"), _t("INR 2.50 lakh (12 months)", (BANK, 1))]])
    memo = _memo(summary=[], sections=[CamSection(id="banking", number=1, title="Banking", blocks=[
        CamBlock(kind="table", table=banking), CamBlock(kind="table", table=facilities)])])
    document = exports.cam_document(memo)
    tables = [b for b in document.blocks if isinstance(b, layout.Table) and b.columns]
    for table, whole in ((tables[0], {0: "Lotuscrest Bank ·1234", 1: "12 (Apr 2025 to Mar 2026)", 3: "INR 52.00 lakh",
                                      5: "INR 60.00 lakh"}),
                         (tables[1], {0: "Tidewater Finance Limited", 1: "Term loan", 2: "INR 80.00 lakh",
                                      3: "INR 95.00 lakh", 4: "INR 95.00 lakh;", 5: "INR 1.45 lakh",
                                      6: "INR 2.50 lakh"})):
        for j, text in whole.items():
            room = table.widths[j] * layout.TEXT_WIDTH - layout.PADDING
            assert stringWidth(layout.winansi(text), "Helvetica", layout.TABLE_SIZE) <= room, (j, text)
    lines = [_n(line) for line in _pages(exports.render(document, "pdf"))[0].splitlines()]
    for text in ("Lotuscrest Bank ·1234", "12 (Apr 2025 to Mar 2026)", "INR 52.00 lakh", "Ashwood Finserv Limited",
                 "Term loan", "INR 80.00 lakh", "INR 95.00 lakh;"):
        assert any(text in line for line in lines), text


def _filler(lines: int) -> CamBlock:
    """A paragraph of about so many lines of body text."""
    return CamBlock(kind="paragraph", paragraph=_t(" ".join(["The statements are summarised below."] * (lines * 3))))


def _rows(count: int, label: str = "Row") -> CamTable:
    return CamTable(columns=["Check", "Outcome"], rows=[[_t(f"{label} {n}"), _t("Passed")] for n in range(1, count + 1)])


def test_a_heading_and_the_line_leading_into_a_table_go_with_the_table_to_the_next_page():
    for lines in range(40, 96):
        memo = _memo(summary=[], sections=[
            CamSection(id="before", number=1, title="Before", blocks=[_filler(lines)]),
            CamSection(id="policy", number=2, title="Compliance with policy norms", blocks=[
                CamBlock(kind="paragraph", paragraph=_t("12 norms apply, on the FY 2025-26 statements.")),
                CamBlock(kind="table", table=_rows(30))])])
        pages = _pages(exports.render(exports.cam_document(memo), "pdf"))
        lead = _page_of(pages, "12 norms apply")
        assert _page_of(pages, "2. Compliance with policy norms") == lead and _row(pages[lead], "Row 1"), lines


def test_a_short_line_after_a_table_is_not_left_alone_at_the_top_of_a_page():
    for lines in range(40, 96):
        memo = _memo(summary=[], sections=[
            CamSection(id="checks", number=1, title="Checks", blocks=[
                _filler(lines), CamBlock(kind="table", table=_rows(4, "Area")),
                CamBlock(kind="paragraph", paragraph=_t("No discrepancy was found."))]),
            CamSection(id="detail", number=2, title="Detail", blocks=[
                CamBlock(kind="table", table=_rows(30)),
                CamBlock(kind="paragraph", paragraph=_t("Every check passed."))])])
        pages = _pages(exports.render(exports.cam_document(memo), "pdf"))
        # A short table goes over whole with the line after it; a long one sends its last row with it.
        short = _page_of(pages, "No discrepancy was found.")
        assert all(_row(pages[short], f"Area {n}") for n in range(1, 5)), lines
        assert _row(pages[_page_of(pages, "Every check passed.")], "Row 30"), lines


def test_a_list_of_sources_starts_under_what_is_before_it_and_leaves_no_row_alone():
    # Fifteen sources, about a third of a page: the note's own seven and eight more its documents cite.
    sources = _note().sources + [_source(n, f"D-{n}", f"Document {n}", f"document-{n}.pdf", [1])
                                 for n in range(8, 16)]
    files = [s.file_name for s in sources]
    split = False
    for count in range(8, 40):
        note = _note(sources=sources, documents_to_collect=[
            _t(f"Document {n} to collect.", (f"D-{n % 8 + 8}", 1)) for n in range(count)])
        data = exports.render(exports.pd_note_document(note), "pdf")
        pages, blank = _pages(data), _blank(data)
        first = _page_of(pages, "application-message.txt")
        # Not alone on a page of its own while the page before has room for its heading and first rows;
        # where it runs over, three rows at least on each page.
        if first and pages[first].lstrip().startswith("Sources"):
            assert blank[first - 1] < 0.2, count
        rows = [sum(f in page for f in files) for page in pages[first:]]
        assert sum(rows) == len(files) and all(r >= 3 for r in rows if r), (count, rows)
        split |= len([r for r in rows if r]) > 1
    assert split
    # Word keeps the first three rows with the heading and the last three together, as the PDF does.
    doc = read_docx(io.BytesIO(exports.render(exports.pd_note_document(_note(sources=sources)), "docx")))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "No." and len(t.columns) == 4)
    kept = [all(p.paragraph_format.keep_with_next for c in row.cells for p in c.paragraphs) for row in table.rows]
    assert kept == [True, True, True] + [False] * (len(sources) - 5) + [True, True, False]


def _blank(data: bytes) -> list[float]:
    """How much of each page's body is left blank under its last mark, as a share of the body's height."""
    scale, top, bottom = 0.5, TOP, 842 - FOOT
    pdf = pdfium.PdfDocument(data)
    try:
        out = []
        for page in pdf:
            image = page.render(scale=scale).to_pil().convert("L")
            body = image.crop((0, round(top * scale), image.width, round(bottom * scale)))
            box = ImageOps.invert(body).point(lambda v: 255 if v > 40 else 0).getbbox()
            out.append(1 - (box[3] if box else 0) / body.height)
        return out
    finally:
        pdf.close()


def _question_with(lines: int, number: int = 1) -> PdQuestion:
    """A question whose reason runs to about so many lines."""
    return _question(number, "turnover", f"Question {number}: where are the rest of the sales receipts banked?",
                     " ".join(["GST outward supplies run above the business credits in the bank."] * (lines * 2)),
                     severity="moderate", source="finding", ref=f"F-{number}", xv_ref=f"XV-0{number}",
                     rule="TO-02", title="GST turnover not supported by bank credits", status="Open",
                     target="cross_verification", evidence=[Evidence(document_id=GST, page=2)])


def test_a_question_that_nearly_fits_turns_the_page_between_its_answer_lines_rather_than_move_whole():
    # A first question long enough to leave a varying room under it for the second.
    turned = False
    for lines in range(6, 30):
        note = _note(people=[], record_fields=[], documents_to_collect=[], sign_off=[],
                     questions=[_question_with(lines), _question_with(2, 2)],
                     topics=[PdTopic(area="turnover", title="Turnover, tax and sales", question_ids=["PDQ-1", "PDQ-2"])])
        data = exports.render(exports.pd_note_document(note), "pdf")
        pages, blank = _pages(data), _blank(data)
        second = _page_of(pages, "Question 2:")
        asked = pages[second].split("Question 2:")[1]
        # Never apart from its reason, its evidence and the label over its response lines.
        assert all(text in asked for text in ("Why we ask: GST", "Evidence: [5 p.2]", "Response")), lines
        # Started on a page with room for it, rather than leave a third of that page blank.
        assert all(b < 0.3 for b in blank[:-1]), (lines, blank)
        turned |= second < len(pages) - 1 and "Follow-up" not in asked
    assert turned
    # Word may turn the page after the second response line, the third, the fourth (before the
    # follow-up), and nowhere else in the question.
    doc = read_docx(io.BytesIO(exports.render(exports.pd_note_document(_note()), "docx")))
    response, follow_up = [t for t in doc.tables if len(t.columns) == 1][:2]
    assert [row.cells[0].paragraphs[0].paragraph_format.keep_with_next for row in response.rows] == [
        True, None, None, None]
    assert [row.cells[0].paragraphs[0].paragraph_format.keep_with_next for row in follow_up.rows] == [True, None]
    paragraphs = doc.paragraphs
    first = next(k for k, p in enumerate(paragraphs) if p.text.startswith("1. "))
    head = paragraphs[first:first + 5]
    assert [p.text.split(":")[0] for p in head] == [paragraphs[first].text, "Serious · Finding XV-01 (ID-01) · Open",
                                                    "Why we ask", "Evidence", "Response"]
    assert all(p.paragraph_format.keep_with_next or p.style.paragraph_format.keep_with_next for p in head)


def test_a_table_of_up_to_twelve_facts_is_kept_on_one_page_with_its_heading():
    facts = [CamFact(label=f"Fact {n}", value=_t(f"Value {n}", (COI, 1))) for n in range(1, 13)]
    for lines in range(40, 96, 3):
        memo = _memo(summary=[], sections=[CamSection(id="policy", number=1, title="Compliance with policy norms",
                                                      blocks=[_filler(lines), CamBlock(kind="subheading",
                                                                                       text="Approving authority"),
                                                              CamBlock(kind="facts", facts=facts)])])
        pages = _pages(exports.render(exports.cam_document(memo), "pdf"))
        here = _page_of(pages, "Approving authority")
        assert all(_row(pages[here], f"Fact {n}") for n in range(1, 13)), lines
    doc = read_docx(io.BytesIO(exports.render(exports.cam_document(memo), "docx")))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Fact 1")
    assert [all(p.paragraph_format.keep_with_next for c in row.cells for p in c.paragraphs) for row in table.rows] \
        == [True] * 11 + [False]


def test_a_long_table_starts_where_it_is_reached_and_repeats_its_header_on_each_page_it_runs_onto():
    # Rows of two to eight lines: the space left under a page's last whole row varies.
    detail = "Credits received from customers, after transfers between own accounts are removed. "
    rows = [[_t(f"Row {n}"), _t(detail * (n % 7 + 2)), _t(f"INR {n}.00 cr", (AFS, 4))] for n in range(1, 61)]
    memo = _memo(sections=[CamSection(id="banking", number=1, title="Banking and account conduct", blocks=[
        CamBlock(kind="paragraph", paragraph=_t("The statements are summarised below. " * 30)),
        CamBlock(kind="table", table=CamTable(title="Credits by month", columns=["Month", "Detail", "Credits"],
                                              numeric=[2], rows=rows))])])
    document = exports.cam_document(memo)
    pages = _pages(exports.render(document, "pdf"))
    first = next(i for i, page in enumerate(pages) if "Credits by month" in page)
    # Under the text before it, not moved to the next page; the header once on each page, at its top.
    assert "The statements are summarised below." in pages[first] and re.search(r"Row 1\b", pages[first])
    running = [page for page in pages if re.search(r"Row \d+\b", page)]
    assert len(running) >= 3 and all(page.count("Month") == 1 for page in running)
    assert all(page.index("Month") < page.index("Row ") for page in running[1:])
    # Word moves a row to the next page whole, as the PDF does.
    doc = read_docx(io.BytesIO(exports.render(document, "docx")))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Month")
    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in table.rows[1:])


def test_the_sources_start_on_their_own_page_after_the_recommendation_and_the_sign_off_keeps_its_page():
    document = exports.cam_document(_memo())
    heading = next(p for p in read_docx(io.BytesIO(exports.render(document, "docx"))).paragraphs if p.text == "Sources")
    assert heading.paragraph_format.page_break_before
    pages = _pages(exports.render(document, "pdf"))
    assert pages[-1].lstrip().startswith("Sources") and "Approved by" in pages[-2]
    # A table kept with its group, as the note's sign-off is, stays on the page where the group fits.
    pages = _pages(exports.render(exports.pd_note_document(_note()), "pdf"))
    assert any("Documents to collect at the discussion" in page and "Discussion held by" in page for page in pages)


def test_the_recommendation_follows_the_last_lines_of_the_summary_where_it_fits_whole_under_them():
    # Open queries enough to end the summary anywhere on a page: the recommendation and its sign-off
    # stay on one page, a page of their own unless they fit whole under the summary's last lines.
    shared = False
    for count in range(1, 70, 3):
        memo = _memo(summary=[CamSummaryItem(kind="query", ref=f"Q-{n}", group="documents_needed",
                                             text=f"GSTR-3B for month {n}.") for n in range(1, count + 1)])
        data = exports.render(exports.cam_document(memo), "pdf")
        pages, blank = _pages(data), _blank(data)
        here = _page_of(pages, "5. Recommendation")
        assert all(prompt in pages[here] for prompt in PROMPTS) and "Approved by" in pages[here], count
        if pages[here].lstrip().startswith("5. Recommendation"):
            # On a page of its own only where the page before has no room for it.
            assert blank[here - 1] < 1 - blank[here] + 0.03, (count, blank)
        else:
            shared = True
    assert shared
    # Word keeps it together, rather than start it on a page of its own whatever is above it.
    doc = read_docx(io.BytesIO(exports.render(exports.cam_document(_memo()), "docx")))
    heading = next(p for p in doc.paragraphs if p.text == "5. Recommendation")
    assert not heading.paragraph_format.page_break_before
    paragraphs = doc.paragraphs
    first = next(k for k, p in enumerate(paragraphs) if p.text == "5. Recommendation")
    kept = paragraphs[first:first + 1 + len(PROMPTS)]
    assert [p.text for p in kept] == ["5. Recommendation", *PROMPTS]
    assert all(p.paragraph_format.keep_with_next for p in kept)


def test_the_same_memo_gives_the_same_files_dated_as_of():
    document = exports.cam_document(_memo())
    for fmt in ("docx", "pdf"):
        assert exports.render(document, fmt) == exports.render(exports.cam_document(_memo()), fmt), fmt
    data = exports.render(document, "docx")
    props = read_docx(io.BytesIO(data)).core_properties
    assert (props.title, props.subject, props.author) == (
        "Credit Approval Memo", "Kestrel Fabricators Pvt Ltd · BBG-2026-000004", "CreditIQ")
    assert props.created.replace(tzinfo=None) == props.modified.replace(tzinfo=None) == datetime(2026, 9, 29)
    assert {i.date_time for i in zipfile.ZipFile(io.BytesIO(data)).infolist()} == {(2026, 9, 29, 0, 0, 0)}
    assert b"/CreationDate (D:20260929000000+00'00')" in exports.render(document, "pdf")


def test_a_borrower_name_longer_than_word_allows_in_a_property_still_gives_both_files():
    borrower = ("Kestrel Precision Engineering and Allied Fabrication Works " * 6)[:300]
    document = exports.cam_document(_memo(borrower=borrower))
    data = exports.render(document, "docx")
    doc = read_docx(io.BytesIO(data))
    assert len(doc.core_properties.subject) == 255 and doc.core_properties.subject.endswith("…")
    # The running header on one line in both files: smaller, then shortened.
    header = doc.sections[0].header.paragraphs[0]
    assert header.text == layout.fit(document.header, 8)[0] and header.text.endswith("…")
    assert header.runs[0].font.size.pt == 6.5
    assert borrower in _texts(_body(data))
    assert _pdf_body(exports.render(document, "pdf"), document)


# --- The downloads ---

def _processed_case(client: TestClient, root: Path) -> tuple[dict, str]:
    analyst = sign_in(client, ANALYST)
    case_id = make_case(client, analyst, root, borrower="Kestrel Fabricators Pvt Ltd", facilities=["cash_credit"],
                        amount_inr=30_000_000)
    assert upload(client, analyst, case_id, case_files("complete")).status_code == 202
    run_jobs(root)
    conn = db.connect(root)
    try:
        pipeline.refresh_case(conn, root, case_id)
    finally:
        conn.close()
    return analyst, case_id


def _download(client: TestClient, headers: dict, case_id: str, path: str, fmt: str):
    return client.get(f"/api/cases/{case_id}/{path}/download", params={"format": fmt}, headers=headers)


def _one_page_statement(memo: Cam) -> tuple[str, Evidence]:
    """A statement of the memo that cites one page of one document."""
    for section in memo.sections:
        for block in section.blocks:
            texts = [block.paragraph] if block.paragraph else [f.value for f in block.facts]
            for t in texts:
                if t and len({(c.document_id, c.page) for c in t.citations}) == 1 and layout.clean(t.text):
                    return layout.clean(t.text), t.citations[0]
    raise AssertionError("no statement cites a single page")


def test_tc26_tc27_a_processed_case_downloads_as_word_and_pdf(client: TestClient, published_data_root: Path):
    analyst, case_id = _processed_case(client, published_data_root)
    memo = Cam.model_validate(client.get(f"/api/cases/{case_id}/cam", headers=analyst).json())
    note = PdNote.model_validate(client.get(f"/api/cases/{case_id}/pd-note", headers=analyst).json())
    files = {}
    # From the screens' origin, as a browser asks: the file name stays readable through the CORS middleware.
    screens = {**analyst, "Origin": SCREENS}
    for output, path, name in (("cam", "cam", "CAM-draft"), ("pd_note", "pd-note", "PD-note")):
        for fmt, media in (("docx", DOCX), ("pdf", "application/pdf")):
            response = _download(client, screens, case_id, path, fmt)
            assert response.status_code == 200, response.text
            assert response.headers["content-type"] == media
            assert response.headers["content-disposition"] == f'attachment; filename="{name}-{case_id}.{fmt}"'
            assert response.headers["access-control-allow-origin"] == SCREENS
            assert response.headers["access-control-expose-headers"] == "Content-Disposition"
            files[output, fmt] = response.content

    body = _body(files["cam", "docx"])
    assert body[0] == ("p", "Memo title", memo.title)
    sources_heading = layout.clean(memo.labels.get("sources")) or "Sources"
    assert _headings(body) == [f"{s.number}. {layout.clean(s.title)}" for s in memo.sections] + [
        f"{memo.summary_number}. {memo.summary_title}", f"{memo.recommendation_number}. {memo.recommendation_title}",
        sources_heading]
    assert memo.recommendation == "" and memo.recommendation_prompts
    for prompt in memo.recommendation_prompts:
        assert _after(body, prompt) == ("t", [[""]]), prompt
    sign_off = next(b[1] for b in body if b[0] == "t" and len(b[1][0]) == 5 and b[1][0][0] == "" and b[1][1][1:] == [""] * 4)
    assert [r[0] for r in sign_off[1:]] == [s.role for s in memo.sign_off]
    assert memo.sources and [int(r[0]) for r in _after(body, sources_heading)[1][1:]] == [s.number for s in memo.sources]
    # One statement's marker names its document by the number the memo's sources give it.
    text, cited = _one_page_statement(memo)
    number = next(s.number for s in memo.sources if s.document_id == cited.document_id)
    assert _n(f"{text} [{number} p.{cited.page}]") in _texts(body)

    assert "page 1 of" in _pages(files["cam", "pdf"])[0]
    pdf = _pdf_body(files["cam", "pdf"], exports.cam_document(memo))
    for section in memo.sections:
        assert _squash(layout.winansi(f"{section.number}. {layout.clean(section.title)}")) in pdf
        for block in section.blocks:
            if block.kind == "paragraph":
                assert _squash(layout.winansi(layout.clean(block.paragraph.text)))[:60] in pdf, block.paragraph.text
    for item in memo.summary:
        # A finding and a check not completed by their titles (the register), the rest by what they say.
        said = item.title if item.kind in ("finding", "incomplete") and item.title else item.text
        assert _squash(layout.winansi(layout.clean(said)))[:60] in pdf, said

    body = _body(files["pd_note", "docx"])
    asked = [b[2] for b in body if b[0] == "p" and b[1] == "Memo question"]
    order = [q for t in note.topics for q in note.questions if q.id in t.question_ids]
    assert asked == [f"{q.number}. {layout.clean(q.question)}" for q in order] and asked
    response = layout.clean(note.labels.get("response")) or "Response"
    assert sum(1 for b in body if b[0] == "p" and b[2] == response) == len(note.questions)
    assert "page 1 of" in _pages(files["pd_note", "pdf"])[0]
    pdf = _pdf_body(files["pd_note", "pdf"], exports.pd_note_document(note))
    assert all(_squash(layout.winansi(f"{q.number}. {layout.clean(q.question)}")) in pdf for q in note.questions)
    assert _pdf_body(files["pd_note", "pdf"], exports.pd_note_document(note)) == _word_as_pdf(files["pd_note", "docx"])

    conn = db.connect(published_data_root)
    try:
        logged = [(r["actor"], json.loads(r["detail"])) for r in conn.execute(
            "SELECT actor, detail FROM decision_log WHERE case_id = ? AND action = 'output.downloaded' ORDER BY seq",
            (case_id,))]
    finally:
        conn.close()
    assert logged == [("ananya-krishnan", {"output": o, "format": f})
                      for o in ("cam", "pd_note") for f in ("docx", "pdf")]


def test_the_downloads_are_read_as_the_outputs_are(client: TestClient, published_data_root: Path):
    rm = sign_in(client, RM)
    own = make_case(client, rm, published_data_root, borrower="Heronbay Polymers Pvt Ltd")
    analyst = sign_in(client, ANALYST)
    case_id = make_case(client, analyst, published_data_root)
    for path in ("cam", "pd-note"):
        assert _download(client, {}, case_id, path, "pdf").status_code == 401
        assert _download(client, rm, case_id, path, "pdf").status_code == 403
        assert _download(client, rm, own, path, "pdf").status_code == 403
        assert _download(client, sign_in(client, ADMIN), case_id, path, "pdf").status_code == 403
        assert _download(client, analyst, case_id, path, "xlsx").status_code == 422
        assert client.get(f"/api/cases/{case_id}/{path}/download", headers=analyst).status_code == 422
        assert _download(client, analyst, "BBG-2099-999999", path, "pdf").status_code == 404
        assert _download(client, sign_in(client, MANAGER), case_id, path, "docx").status_code == 200
    conn = db.connect(published_data_root)
    try:
        logged = conn.execute("SELECT COUNT(*) FROM decision_log WHERE action = 'output.downloaded'").fetchone()[0]
    finally:
        conn.close()
    assert logged == 2  # only the two that were served


def test_a_role_that_reads_outputs_but_not_every_case_downloads_only_its_own(
        client: TestClient, published_data_root: Path, monkeypatch: pytest.MonkeyPatch):
    # Every configured role that reads outputs reads every case too: an RM given output.read shows the
    # downloads keep to the cases the user may see, as the outputs do.
    permissions = auth.role_permissions
    monkeypatch.setattr(auth, "role_permissions", lambda role: permissions(role) + ["output.read"] * (role == "rm"))
    rm = sign_in(client, RM)
    own = make_case(client, rm, published_data_root, borrower="Heronbay Polymers Pvt Ltd")
    other = make_case(client, sign_in(client, ANALYST), published_data_root)
    for path in ("cam", "pd-note"):
        assert _download(client, rm, own, path, "pdf").status_code == 200
        assert _download(client, rm, other, path, "pdf").status_code == 404
        assert client.get(f"/api/cases/{other}/{path}", headers=rm).status_code == 404


def test_a_file_that_cannot_be_written_is_a_worded_error_the_screens_can_read(
        client: TestClient, published_data_root: Path, monkeypatch: pytest.MonkeyPatch):
    analyst = sign_in(client, ANALYST)
    borrower = ("Heronbay Polymers and Allied Speciality Chemicals " * 7)[:300]
    case_id = make_case(client, analyst, published_data_root, borrower=borrower)
    screens = {**analyst, "Origin": SCREENS}
    for path in ("cam", "pd-note"):
        response = _download(client, screens, case_id, path, "docx")
        assert response.status_code == 200, response.text
        assert len(read_docx(io.BytesIO(response.content)).core_properties.subject) == 255

    def broken(document, fmt):
        raise ValueError("the writer failed")

    monkeypatch.setattr(exports, "render", broken)
    for path, fmt, words in (("cam", "docx", "The draft CAM could not be written as a Word file (ValueError)."),
                             ("pd-note", "pdf", "The PD note could not be written as a PDF file (ValueError).")):
        response = _download(client, screens, case_id, path, fmt)
        assert response.status_code == 500 and response.json() == {"detail": words}
        assert response.headers["access-control-allow-origin"] == SCREENS
    conn = db.connect(published_data_root)
    try:
        logged = conn.execute("SELECT COUNT(*) FROM decision_log WHERE action = 'output.downloaded'").fetchone()[0]
    finally:
        conn.close()
    assert logged == 2  # the two served before the writer failed


def test_the_downloads_need_a_published_configuration(client: TestClient, published_data_root: Path):
    analyst = sign_in(client, ANALYST)
    case_id = make_case(client, analyst, published_data_root)
    (published_data_root / "config_store" / "published.json").unlink()
    for path in ("cam", "pd-note"):
        assert _download(client, analyst, case_id, path, "pdf").status_code == 503
